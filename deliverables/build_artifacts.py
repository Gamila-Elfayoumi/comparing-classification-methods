"""Build the research report and presentation from measured CSV/PNG outputs.

Run with the Codex bundled Python runtime after src/analyze.py.
"""
from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
RESULTS = ROOT / "results"
FIG = RESULTS / "figures"
OUT.mkdir(exist_ok=True)
METHODS = ["PCA", "Kernel PCA", "SNE", "t-SNE", "PCA + t-SNE", "UMAP", "MDS", "PCoA", "Isomap"]
DATASETS = ["HAM10000", "Blood cell"]
NAVY = "153B50"
BLUE = "23668A"
GRAY = "5C6770"

def read_csv(path):
    with path.open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))

summary = {(r["dataset"], r["method"]): r for r in read_csv(RESULTS / "primary_summary.csv")}
param = read_csv(RESULTS / "parameter_comparison.csv")
runs = read_csv(RESULTS / "runs.csv")
paired = read_csv(RESULTS / "paired_tests.csv")

def value(dataset, method, metric, digits=3):
    return f"{float(summary[(dataset, method)][metric + '_mean']):.{digits}f}"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def add_table(doc, headers, body, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = DInches(width)
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = header
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255,255,255)
            run.font.size = DPt(8)
    for i, row in enumerate(body):
        cells = table.add_row().cells
        for j, item in enumerate(row):
            cells[j].text = str(item)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2:
                set_cell_shading(cells[j], "EDF3F6")
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = DPt(0)
                for run in p.runs:
                    run.font.size = DPt(7.8)
    doc.add_paragraph()
    return table

def add_figure(doc, filename, caption, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG / filename), width=DInches(width))
    cap = doc.add_paragraph(caption)
    cap.style = "Caption"
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

def add_metric_tables(doc, dataset):
    doc.add_heading(f"{dataset} measured performance", 2)
    body = []
    for method in METHODS:
        body.append([method, value(dataset, method, "accuracy"),
                     value(dataset, method, "sensitivity"), value(dataset, method, "specificity"),
                     value(dataset, method, "precision"), value(dataset, method, "f_measure"),
                     value(dataset, method, "auc_roc")])
    add_table(doc, ["Method", "Accuracy", "Sensitivity / recall", "Specificity", "Precision", "F1", "ROC AUC"], body,
              [1.25, .75, 1.05, .85, .85, .55, .75])
    body = []
    for method in METHODS:
        body.append([method, value(dataset, method, "error"), value(dataset, method, "kappa"),
                     value(dataset, method, "informedness"), value(dataset, method, "markedness"),
                     value(dataset, method, "log_loss"), value(dataset, method, "cpu_seconds", 2)])
    add_table(doc, ["Method", "Error", "Kappa", "Informedness", "Markedness", "Log loss", "CPU s"], body,
              [1.25, .65, .65, 1.05, 1.05, .8, .65])
    doc.add_paragraph("Values are means across three held-out splits for prespecified setting A. Recall equals sensitivity. CPU time includes embedding, held-out transform, classifier fit, and prediction, after the separately reported UMAP compilation warmup.")

def report():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = DInches(8.5), DInches(11)
    sec.top_margin = sec.bottom_margin = DInches(.72)
    sec.left_margin = sec.right_margin = DInches(.78)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = DPt(10)
    styles["Normal"].paragraph_format.space_after = DPt(6)
    for h, size in [("Title", 23), ("Heading 1", 15), ("Heading 2", 12)]:
        styles[h].font.name = "Aptos Display"
        styles[h].font.size = DPt(size)
        styles[h].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Title"].font.color.rgb = RGBColor(0,0,0)
    title_properties = styles["Title"]._element.get_or_add_pPr()
    for node in list(title_properties):
        if node.tag == qn("w:pBdr"):
            title_properties.remove(node)
    styles["Caption"].font.size = DPt(8)
    styles["Caption"].font.italic = True
    title = doc.add_paragraph(style="Title")
    title.add_run("Dimensionality Reduction for Two Biomedical Classification Tables")
    p = doc.add_paragraph("Measured pilot comparison of nine dimensionality reduction methods")
    p.style = "Subtitle"
    doc.add_paragraph("Research report | 21 October 2026 | Reproducible code and raw run tables accompany this document")
    doc.add_heading("Abstract", 1)
    ham_best = max(METHODS, key=lambda m: float(summary[("HAM10000",m)]["auc_roc_mean"]))
    blood_best = max(METHODS, key=lambda m: float(summary[("Blood cell",m)]["auc_roc_mean"]))
    doc.add_paragraph(
        f"This preliminary study compares nine two-dimensional representations on a lesion image table and a blood-cell feature table. "
        f"Across three shared splits, {ham_best} had the highest mean ROC AUC on HAM10000 ({value('HAM10000',ham_best,'auc_roc')}), "
        f"while {blood_best} had the highest mean ROC AUC on the blood-cell table ({value('Blood cell',blood_best,'auc_roc')}). "
        "Performance varied across splits and parameter settings. The data and validation design do not support clinical or journal-level generalization claims. FROC AUC is unavailable because the files contain no candidate lesion locations.")
    doc.add_heading("Introduction", 1)
    doc.add_paragraph("Dimensionality reduction is often used to inspect high-dimensional biomedical data. A visually separated map does not automatically retain information needed for a new patient or image. This benchmark therefore compares embeddings through a fixed downstream classifier on held-out samples, while measuring runtime and parameter sensitivity. It treats the methods as representation procedures, not standalone classifiers.")
    doc.add_paragraph("PCA is linear and maximizes retained variance. Kernel PCA applies a radial basis function kernel. Gaussian SNE matches neighbor probabilities with a Gaussian low-dimensional affinity. t-SNE uses a heavy-tailed affinity. The hybrid applies PCA before t-SNE. UMAP builds a fuzzy neighborhood graph. Metric MDS preserves pairwise distances, PCoA decomposes a centered distance matrix, and Isomap approximates geodesic distances. All methods output two coordinates for a common comparison.")
    doc.add_paragraph("Each method preserves a different property, so a good two-dimensional picture is not necessarily a good classifier input. PCA keeps directions of greatest overall variance, which may reflect illumination or cell size rather than diagnosis. Kernel PCA can unfold curved structure but depends on kernel scale. MDS and PCoA emphasize global distances; the chosen distance measure determines which differences matter. Isomap estimates distances along a neighbor graph and can help when class information follows a curved manifold, but graph construction can be unstable with sparse or noisy samples. SNE, t-SNE, and UMAP focus on local neighborhoods. Their visual clusters can be useful for exploration, yet cluster separation alone does not establish predictive validity on new records.")
    doc.add_heading("Dataset description", 1)
    add_table(doc, ["Dataset", "Records", "Input features", "Positive class", "Split unit"], [
        ["HAM10000", "10,015", "192 values from 8 × 8 RGB pixels", "Melanoma: 1,113", "Lesion ID"],
        ["Blood cell anomaly", "5,880", "22 morphology, color, and blood-count values", "Anomaly: 1,880", "Row; patient ID absent"],
    ], [1.2,.7,2.1,1.1,1.15])
    doc.add_paragraph("HAM10000 pixel rows were checked against diagnosis metadata, with exact agreement for all 10,015 labels. Lesions are separated between train and test. The 8 × 8 representation reduces compute but discards clinically important detail. The blood-cell table excludes cell type, disease category, anomaly score, classification confidence, and labeller confidence because these fields can directly reveal the target. Its origin and patient independence remain unverified.")
    doc.add_heading("Preprocessing and evaluation", 1)
    doc.add_paragraph("For each of three seeds, the study draws 750 training records and 300 held-out records after a group split for HAM10000 or a stratified row split for the blood-cell table. Subsampling retains approximate class proportions. Standardization fits on the training data only. Each embedding fits on training data only. A balanced logistic regression is trained in two dimensions and evaluated at a fixed 0.5 probability threshold. This design avoids fitting an embedding on the test observations.")
    doc.add_paragraph("PCA, kernel PCA, UMAP, and Isomap transform held-out records directly. PCoA uses a Gower distance extension. The SNE, t-SNE, and MDS implementations lack an inductive transform; a distance-weighted five-neighbor regression approximates test coordinates from the training embedding. This changes what is being evaluated, so rankings should be read as end-to-end pipeline rankings.")
    add_table(doc, ["Method", "Setting A", "Setting B"], [
        ["PCA", "No whitening", "Whitening"], ["Kernel PCA", "RBF gamma = 0.5 / p", "RBF gamma = 2 / p"],
        ["SNE", "Perplexity 15", "Perplexity 30"], ["t-SNE", "Perplexity 15", "Perplexity 30"],
        ["PCA + t-SNE", "PCA to ≤30 then perplexity 15", "PCA to ≤30 then perplexity 30"],
        ["UMAP", "15 neighbors", "30 neighbors"],
        ["Metric MDS", "80 iterations", "150 iterations"],
        ["PCoA", "Cosine distance", "Manhattan distance"],
        ["Isomap", "10 neighbors", "25 neighbors"],
    ], [1.3,2.25,2.25])
    doc.add_heading("Metric definitions", 1)
    doc.add_paragraph("From true positives TP, true negatives TN, false positives FP, and false negatives FN: accuracy = (TP+TN)/N; error = 1−accuracy; sensitivity or recall = TP/(TP+FN); specificity = TN/(TN+FP); precision = TP/(TP+FP); F1 = 2·precision·recall/(precision+recall). ROC AUC integrates sensitivity against false-positive rate across thresholds. Kappa is chance-corrected agreement. Informedness = sensitivity+specificity−1. Markedness = precision+negative predictive value−1. Log loss evaluates probability calibration and discrimination together.")
    doc.add_paragraph("PCoA is defined by its distance matrix. Here cosine distance is primary and Manhattan distance is the sensitivity setting; Euclidean PCoA would largely duplicate PCA. MDS and Isomap use Euclidean input distances. Isomap's neighbor count controls its graph connectivity.")
    doc.add_heading("Results and discussion", 1)
    for dataset in DATASETS:
        add_metric_tables(doc, dataset)
        best = max(METHODS, key=lambda m: float(summary[(dataset,m)]["auc_roc_mean"]))
        low = min(METHODS, key=lambda m: float(summary[(dataset,m)]["cpu_seconds_mean"]))
        doc.add_paragraph(f"In {dataset}, {best} has the highest mean ROC AUC for setting A ({value(dataset,best,'auc_roc')}). {low} uses the least mean CPU time ({value(dataset,low,'cpu_seconds',2)} s). These are descriptive rankings across only three overlapping splits. Accuracy should be read together with sensitivity and specificity because the positive class is less frequent.")
        if dataset == "HAM10000":
            doc.add_paragraph("Isomap's mean AUC of 0.649 exceeds PCA's 0.602, suggesting that a neighborhood-graph representation may retain more melanoma ranking information in these compressed pixels. This is a hypothesis about the observed data, not proof of an underlying disease manifold. PCA still has slightly higher accuracy (0.609 versus 0.590), showing that ranking quality and decisions at the fixed 0.5 threshold answer different questions. Isomap's precision is only 0.134: many positive predictions are false alarms in this low-prevalence sample. Neither AUC nor accuracy is strong enough here to support a diagnostic claim.")
        else:
            doc.add_paragraph("PCA has the highest mean AUC on the blood-cell table (0.630) while using about 0.02 CPU seconds, consistent with much of the usable signal being captured by broad linear variation in these selected numeric features. SNE has higher accuracy (0.614) but lower AUC (0.581), because one fixed decision threshold can favor the majority class without ranking anomalous records well. MDS has the lowest AUC (0.562) and high CPU cost (about 1.06 seconds); preserving global Euclidean distances in two coordinates did not preserve the class signal well in this pilot. These explanations remain plausible interpretations rather than established mechanisms.")
        stem = dataset.replace(" ", "_")
        add_figure(doc, f"{stem}_accuracy.png", f"Figure. Held-out accuracy by method on {dataset}.")
        add_figure(doc, f"{stem}_loss.png", f"Figure. Common held-out classifier log loss on {dataset}.")
        add_figure(doc, f"{stem}_roc.png", f"Figure. ROC curves for the first shared split on {dataset}; legend gives split-specific AUC.")
        add_figure(doc, f"{stem}_confusion.png", f"Figure. Held-out confusion matrices for the first shared split on {dataset}.")
        add_figure(doc, f"{stem}_cpu.png", f"Figure. Mean process CPU time per run on {dataset}, excluding UMAP compilation warmup.")
        add_figure(doc, f"{stem}_parameters.png", f"Figure. Mean ROC AUC for two prespecified parameter settings on {dataset}.")
    doc.add_heading("Optimization loss and FROC", 1)
    add_figure(doc, "sne_objective.png", "Figure. Gaussian SNE optimization objective for the first split and setting A. This loss is not directly comparable with the common classifier log loss.")
    doc.add_paragraph("The UMAP and PCA objectives have different units and meanings, so their native optimization losses are not plotted on one axis. The common held-out logistic log loss above is the meaningful cross-method loss comparison.")
    add_figure(doc, "froc_applicability.png", "Figure. Why a valid FROC AUC cannot be calculated from the supplied tables.")
    doc.add_paragraph("A free-response operating characteristic requires lesion locations and scored candidate detections per image. The available files provide one class label and one classifier score per record. Counting each positive classification as a detection would simply relabel ROC and falsely imply localization performance. FROC AUC is therefore reported as not estimable.")
    doc.add_heading("Statistical comparisons", 1)
    doc.add_paragraph("The repository applies paired t and Wilcoxon signed-rank comparisons of AUC against PCA across three shared splits, Holm correction for paired t tests, repeated-measures ANOVA across methods, a requested descriptive MANOVA of AUC and log loss, and McNemar's discordance z statistic on the first split. The paired analyses respect split matching. The repeated splits overlap in training observations and provide weak inferential power; p values are exploratory. MANOVA's independence assumption is violated by repeated splits, so its output is retained as a diagnostic only.")
    body=[]
    for r in paired:
        body.append([r["dataset"], r["comparison"], f"{float(r['mean_difference']):+.3f}",
                     f"{float(r['paired_t_holm_p']):.3f}", f"{float(r['wilcoxon_p']):.3f}"])
    add_table(doc, ["Dataset", "AUC comparison", "Mean Δ", "Paired t Holm p", "Wilcoxon p"], body,
              [1.05,1.7,.65,1.05,.9])
    doc.add_paragraph("The repeated-measures ANOVA of ROC AUC gives F(8,40)=0.673 and p=0.712. The descriptive MANOVA method term gives Wilks' lambda=0.840 and p=0.945, but its independence assumption is violated. On the first split, McNemar z for UMAP versus PCA is 0.667 (p=0.505) on HAM10000 and −2.201 (uncorrected p=0.028) on blood cells; the latter is one of eight exploratory contrasts and does not survive Bonferroni correction. Full outputs are in the results directory. These tests do not substitute for nested cross-validation, independent patient cohorts, confidence intervals, and a prospective analysis plan.")
    doc.add_heading("Limitations and next study", 1)
    doc.add_paragraph("The study compresses lesion images to 8 × 8 pixels and forces every representation to two dimensions. It has three repeated splits, a small subsample, an approximate t-SNE/SNE test transform, and no independent blood-cell patient identifier. The blood-cell source's clinical provenance is not verified. No external cohort, fairness assessment, or uncertainty interval is available. These limits prevent a clinical diagnostic claim or any promise of Q1 journal acceptance. A publication study should use original images or validated embeddings, prespecified nested validation, external cohorts, and lesion detection annotations if FROC is a target.")
    doc.add_heading("References and data sources", 1)
    refs = [
        "HAM10000 dataset: https://www.kaggle.com/datasets/kmader/skin-cancer-mnist-ham10000",
        "HAM10000 source paper: https://arxiv.org/abs/1803.10417",
        "Blood-cell data: https://www.kaggle.com/datasets/alitaqishah/blood-cell-anomaly-detection-2025",
        "Kernel PCA: https://is.mpg.de/de/publications/730",
        "SNE: https://papers.nips.cc/paper/2276-stochastic-neighbor-embedding",
        "t-SNE: https://www.jmlr.org/papers/v9/vandermaaten08a.html",
        "UMAP: https://arxiv.org/abs/1802.03426",
        "Metric MDS implementation: https://scikit-learn.org/stable/modules/generated/sklearn.manifold.MDS.html",
        "Principal coordinates analysis: https://doi.org/10.1093/biomet/53.3-4.325",
        "Isomap: https://doi.org/10.1126/science.290.5500.2319",
    ]
    for ref in refs:
        doc.add_paragraph(ref, style="List Bullet")
    footer = sec.footer.paragraphs[0]
    footer.text = "Preliminary computational study | Reproduce from the accompanying repository"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(OUT / "AML_research_report.docx")

def add_text(slide, text, left, top, width, height, size=18, bold=False, color=PRGB(21,59,80)):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box

def add_slide(prs, title, body=None, image=None, image2=None, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, title, .55, .25, 12.0, .7, 29, True)
    if image and image2:
        slide.shapes.add_picture(str(FIG/image), Inches(.55), Inches(1.15), width=Inches(6.1))
        slide.shapes.add_picture(str(FIG/image2), Inches(6.8), Inches(1.15), width=Inches(6.0))
        if body:
            add_text(slide, body, .75, 5.35, 11.9, 1.15, 17)
    elif image:
        slide.shapes.add_picture(str(FIG/image), Inches(.65), Inches(1.2), width=Inches(8.0))
        if body:
            add_text(slide, body, 8.85, 1.45, 3.85, 4.9, 18)
    elif body:
        add_text(slide, body, .8, 1.35, 11.8, 5.3, 21)
    if footer:
        add_text(slide, footer, .65, 7.05, 12.0, .3, 10, False, PRGB(92,103,112))
    return slide

def slides():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "Dimensionality reduction for biomedical classification", .75, 1.0, 11.8, 1.5, 36, True)
    add_text(slide, "HAM10000 and blood-cell anomaly data\nMeasured pilot comparison | 21 October 2026", .78, 3.0, 11.0, 1.6, 22)
    add_slide(prs, "Study question and design",
              "Can a two-dimensional representation preserve class information for held-out records?\n\nNine algorithms. Two parameter settings each. Three shared splits. One balanced logistic regression per embedding.\n\nPrimary ranking uses setting A only; setting B tests parameter sensitivity.")
    add_slide(prs, "Two datasets and their limits",
              "HAM10000: 10,015 8 × 8 RGB lesion records, 192 pixel features, melanoma versus other diagnoses. Split by lesion ID.\n\nBlood-cell anomaly table: 5,880 rows, 22 selected numeric features, anomaly versus normal. Patient IDs and independent clinical provenance are unavailable.\n\nNeither dataset supports diagnostic deployment from this pilot.")
    add_slide(prs, "What the nine methods do",
              "PCA and kernel PCA retain linear or kernel structure. SNE, t-SNE, the PCA plus t-SNE hybrid, and UMAP emphasize neighborhoods.\n\nMetric MDS preserves Euclidean distances. PCoA decomposes cosine distances. Isomap uses graph geodesics.\n\nAll output two coordinates, then use the same classifier.")
    add_slide(prs, "The comparison measures retained class information",
              "Input features are standardized using the training records. Each algorithm reduces them to two coordinates. The same balanced logistic regression then predicts held-out labels.\n\nAUC asks whether positives generally receive higher scores than negatives across thresholds. Accuracy evaluates one fixed threshold. Log loss penalizes confident wrong probabilities. CPU time covers the full fitted pipeline.\n\nA visually neat map is insufficient evidence unless held-out prediction also improves.")
    add_slide(prs, "Global geometry: PCA, kernel PCA, MDS, PCoA",
              "PCA keeps directions with the most variance. It is fast, but large variance need not represent disease.\n\nKernel PCA uses an RBF similarity so curved patterns can become separable. Its gamma setting determines the scale of similarity.\n\nMDS tries to preserve Euclidean distances between every pair. PCoA eigendecomposes a centered distance matrix; here cosine distance is primary. Both can devote two dimensions to variation unrelated to the label.")
    add_slide(prs, "Local geometry: SNE, t-SNE, hybrid, UMAP, Isomap",
              "SNE and t-SNE match neighbor probabilities. t-SNE's heavy tail reduces crowding; perplexity controls neighborhood scale. PCA plus t-SNE first filters the inputs to at most 30 components.\n\nUMAP builds a fuzzy neighbor graph; its neighbor count controls how local the map is. Isomap uses graph shortest paths to estimate geodesic distances.\n\nThese maps can reveal local structure, but they may distort global distances and change when the neighborhood parameter changes.")
    add_slide(prs, "How unseen test records enter each map",
              "PCA, kernel PCA, UMAP, and Isomap transform new records directly. PCoA uses a distance-based Gower extension.\n\nThe chosen SNE, t-SNE, and MDS implementations have no native test transform. A five-neighbor weighted regressor estimates test coordinates from their training maps.\n\nThe measured ranking therefore compares complete representation pipelines, including this approximation. It does not compare only the training maps.")
    add_slide(prs, "Method readings: PCA through the hybrid",
              "PCA: blood-cell AUC 0.630, the best there. A simple linear map may suffice for these selected features; it is weakest on HAM AUC at 0.602.\n\nKernel PCA: 0.638 HAM, 0.624 blood. Curved RBF similarities help modestly, with kernel scale affecting the map.\n\nSNE: 0.616 HAM, 0.581 blood. Local-neighbor matching does not consistently retain the label signal.\n\nt-SNE: 0.638 HAM, 0.622 blood. Better than SNE on both, but its approximate test transform limits interpretation.\n\nPCA plus t-SNE: 0.632 HAM, 0.616 blood. The preprocessing step shows no gain over t-SNE in this pilot.")
    add_slide(prs, "Method readings: UMAP through Isomap",
              "UMAP: 0.645 HAM, 0.600 blood. The neighbor graph ranks lesions well, but costs about 2.24 CPU seconds per HAM run.\n\nMDS: 0.620 HAM, 0.562 blood. Preserving global Euclidean distances gives the weakest blood-cell AUC.\n\nPCoA: 0.633 HAM, 0.616 blood. Cosine distances give middle-ranking performance; a different distance changes the result.\n\nIsomap: 0.649 HAM, 0.582 blood. Graph geodesics lead HAM AUC but transfer poorly to the blood-cell table.\n\nThe mechanisms above are interpretations of measured associations, not proven causes.")
    for dataset in DATASETS:
        stem=dataset.replace(" ","_")
        best=max(METHODS,key=lambda m:float(summary[(dataset,m)]["auc_roc_mean"]))
        add_slide(prs, f"{dataset}: accuracy and loss",
                  f"Highest mean AUC: {best}, {value(dataset,best,'auc_roc')}. The two charts show why accuracy and probability loss should be read together.",
                  f"{stem}_accuracy.png", f"{stem}_loss.png")
        add_slide(prs, f"{dataset}: ROC and confusion matrices",
                  None, f"{stem}_roc.png", f"{stem}_confusion.png")
        if dataset == "HAM10000":
            add_slide(prs, "HAM10000: better ranking, weak positive precision",
                      "Isomap leads mean ROC AUC at 0.649, versus 0.602 for PCA. A neighbor graph may retain class-related structure that the largest-variance PCA axes miss. The experiment does not establish that mechanism.\n\nPCA has slightly higher accuracy, 0.609 versus 0.590. AUC measures ranking across thresholds, while accuracy uses the prespecified 0.5 threshold.\n\nIsomap precision is 0.134: most positive calls are false positives. The 8 × 8 pixels and low melanoma prevalence limit clinical meaning.")
        else:
            add_slide(prs, "Blood cells: fast PCA wins on ranking",
                      "PCA leads mean ROC AUC at 0.630 and takes about 0.02 CPU seconds. Broad linear variation may capture useful signal in the 22 selected features.\n\nSNE reaches 0.614 accuracy but only 0.581 AUC. Its threshold decisions look better than its overall ranking; accuracy alone would give a misleading winner.\n\nMDS has 0.562 AUC and takes about 1.06 CPU seconds. Keeping global distances in two dimensions did not retain much class signal here. Patient-level generalization remains untested.")
    add_slide(prs, "CPU time has a steep method cost",
              "PCA and kernel PCA are fastest in this sample-size range. UMAP, SNE, and t-SNE require substantially more CPU time. UMAP's first compilation is measured separately.",
              "HAM10000_cpu.png", "Blood_cell_cpu.png")
    add_slide(prs, "CPU cost changes the practical choice",
              "On blood cells, PCA needs about 0.02 CPU seconds and SNE about 6.10 seconds per run. SNE does not gain AUC in exchange for that extra compute (0.581 versus 0.630).\n\nOn HAM10000, Isomap leads AUC at about 0.10 CPU seconds; UMAP is close in AUC (0.645 versus 0.649) but takes about 2.24 seconds.\n\nThese are process CPU times for 750 training and 300 test records. UMAP's one-time compilation warmup is excluded and recorded separately; scaling can change the order.")
    add_slide(prs, "Parameter settings change the ranking",
              "A second setting changes neighborhood size or perplexity. Selecting a winner on the test set would inflate the reported result.",
              "HAM10000_parameters.png", "Blood_cell_parameters.png")
    add_slide(prs, "What the parameter comparison can and cannot say",
              "Setting A was declared primary before reading test results. Setting B checks whether whitening, kernel scale, perplexity, neighbor count, MDS iterations, or PCoA distance changes the result.\n\nLarge A-to-B shifts show tuning sensitivity, especially for methods that define neighborhoods. They do not justify choosing the better test-set setting after the fact.\n\nA stronger study would tune parameters inside each training fold and then evaluate once on an untouched outer test fold.")
    add_slide(prs, "Gaussian SNE optimization loss",
              "The plotted KL objective tracks Gaussian SNE fitting. Native PCA, t-SNE, and UMAP objectives have different meanings. The earlier classifier log-loss charts are comparable across methods.",
              "sne_objective.png")
    add_slide(prs, "FROC AUC is not available",
              "ROC AUC is reported. A valid FROC needs scored candidate lesion locations and false positives per image. These tables have one label and one score per record.",
              "froc_applicability.png")
    add_slide(prs, "Statistical tests are exploratory",
              "Repeated-measures ANOVA of AUC: F(8,40)=0.673, p=0.712. Descriptive MANOVA method term: Wilks' lambda=0.840, p=0.945.\n\nPaired t, Wilcoxon, and McNemar z comparisons are in the repository. Three overlapping splits have low power. MANOVA's independence assumption is violated. Read all p values as diagnostics.")
    add_slide(prs, "What counts as good or bad evidence here",
              "A good result would combine high held-out AUC with useful sensitivity and precision, low log loss, stable parameters, and an acceptable CPU cost. None of these pilot AUCs exceeds 0.65.\n\nA method can look good on one metric and weak on another: SNE has the highest blood-cell accuracy but lower AUC than PCA; Isomap leads HAM AUC but its positive precision is only 0.134.\n\nThree overlapping splits do not provide a reliable superiority claim. The statistical tests and plots describe this experiment, not clinical readiness.")
    add_slide(prs, "Interpretation and next steps",
              "Performance changes with split and parameter choice. The 8 × 8 lesion images limit clinical meaning. The blood-cell table lacks patient-level validation.\n\nFor a journal study: use original or validated features, external cohorts, nested tuning, uncertainty intervals, and lesion locations if FROC is required.",
              footer="Sources and exact numbers: AML_research_report.docx and results/primary_summary.csv")
    prs.save(OUT / "AML_presentation.pptx")

if __name__ == "__main__":
    report()
    slides()
